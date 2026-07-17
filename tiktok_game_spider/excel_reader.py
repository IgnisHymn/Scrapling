"""Excel文件读取模块"""
import openpyxl
from pathlib import Path


class ExcelReader:
    """读取Excel文件中的小游戏配置"""
    
    def __init__(self, excel_path):
        self.excel_path = Path(excel_path)
    
    def read_apps(self):
        """读取AppInfos sheet中的游戏列表"""
        wb = openpyxl.load_workbook(self.excel_path, read_only=True)
        
        # 读取Sheet1: AppInfos
        ws = wb["AppInfos"]
        apps = []
        for row in ws.iter_rows(min_row=1, values_only=True):
            if row[0] and row[1]:  # 确保名称和app_id都存在
                apps.append({
                    "name": str(row[0]).strip(),
                    "app_id": str(row[1]).strip()
                })
        
        wb.close()
        return apps
    
    def read_base_url(self):
        """读取Link sheet中的基础URL"""
        wb = openpyxl.load_workbook(self.excel_path, read_only=True)
        
        # 读取Sheet2: Link
        ws = wb["Link"]
        base_url = ws["A1"].value
        
        wb.close()
        return base_url.strip() if base_url else "https://developers.tiktok.com/portal/game"
