"""Excel报告保存模块"""
import openpyxl
from openpyxl.styles import Alignment, PatternFill, Font
from pathlib import Path


class ExcelSaver:
    """将爬取数据保存到Excel文件"""
    
    # 新模板表头（18列）
    HEADER_ROW = [
        'Date',
        'Ad requests', 'Ad impressions', 'Ad clicks', 'Ad click-through rate', 'eCPM', 'Ad revenue',
        'Total users', 'New users', 'Active users', 'Repeat users',
        'Launched sessions', 'Average launched sessions', 'Average duration per user', 'Average duration per session',
        'Launch success rate', 'Average first-time launch speed', 'Average launch speed'
    ]
    
    # 列映射：字段名 -> 列索引 (0-based)
    COLUMN_MAP = {
        'Date': 0,
        'Ad requests': 1, 'Ad impressions': 2, 'Ad clicks': 3, 'Ad click-through rate': 4, 'eCPM': 5, 'Ad revenue': 6,
        'Total users': 7, 'New users': 8, 'Active users': 9, 'Repeat users': 10,
        'Launched sessions': 11, 'Average launched sessions': 12, 'Average duration per user': 13, 'Average duration per session': 14,
        'Launch success rate': 15, 'Average first-time launch speed': 16, 'Average launch speed': 17
    }
    
    def __init__(self, excel_path):
        self.excel_path = Path(excel_path)
    
    def _find_existing_row(self, ws, date_str):
        """查找指定日期的数据行"""
        for row_idx in range(2, ws.max_row + 1):
            cell_value = ws.cell(row=row_idx, column=1).value
            if cell_value is not None and str(cell_value).strip() == date_str:
                return row_idx
        return None
    
    def _sanitize_sheet_name(self, name):
        """清理sheet名称，移除非法字符"""
        illegal_chars = [':', '\\', '/', '?', '*', '[', ']']
        for char in illegal_chars:
            name = name.replace(char, '_')
        return name[:31]
    
    def _build_row_data(self, date_str, monetization_data, dashboard_data):
        """构建数据行（18列）"""
        row_data = [''] * 18
        
        # Date
        row_data[0] = date_str
        
        # 变现数据
        for field, value in monetization_data.items():
            if field in self.COLUMN_MAP:
                col_idx = self.COLUMN_MAP[field]
                row_data[col_idx] = value
        
        # 仪表板数据（Users + Performance）
        for field, value in dashboard_data.items():
            if field in self.COLUMN_MAP:
                col_idx = self.COLUMN_MAP[field]
                row_data[col_idx] = value
        
        return row_data
    
    def save_data(self, app_name, date_str, monetization_data, dashboard_data):
        """保存数据到Excel（支持去重和更新）"""
        # 加载或创建工作簿
        if self.excel_path.exists():
            wb = openpyxl.load_workbook(self.excel_path)
        else:
            wb = openpyxl.Workbook()
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]
        
        # 获取或创建sheet
        sheet_name = self._sanitize_sheet_name(app_name)
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(title=sheet_name)
            self._init_header(ws)
        
        # 查找已存在的数据行
        existing_row = self._find_existing_row(ws, date_str)
        
        # 构建数据行
        row_data = self._build_row_data(date_str, monetization_data, dashboard_data)
        
        # 写入数据
        target_row = existing_row if existing_row else ws.max_row + 1
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=target_row, column=col_idx, value=value)
        
        # 保存
        wb.save(self.excel_path)
        
        action = "更新" if existing_row else "追加"
        print(f"  数据已{action}到 {sheet_name} sheet 第 {target_row} 行")
    
    def _init_header(self, ws):
        """初始化表头（简化版）"""
        # 写入表头
        for col_idx, value in enumerate(self.HEADER_ROW, start=1):
            cell = ws.cell(row=1, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.font = Font(name="Arial", size=10, bold=True, color="1F2329")
        
        # 背景颜色
        fill_monetization = PatternFill(start_color="E1EAFF", end_color="E1EAFF", fill_type="solid")  # 淡蓝
        fill_users = PatternFill(start_color="FBBFBC", end_color="FBBFBC", fill_type="solid")  # 淡粉红
        fill_performance = PatternFill(start_color="D9F5D6", end_color="D9F5D6", fill_type="solid")  # 淡绿
        
        # 填充背景色
        for col_idx in range(1, 19):
            cell = ws.cell(row=1, column=col_idx)
            if 2 <= col_idx <= 7:  # 变现数据
                cell.fill = fill_monetization
            elif 8 <= col_idx <= 15:  # Users数据
                cell.fill = fill_users
            elif 16 <= col_idx <= 18:  # Performance数据
                cell.fill = fill_performance
